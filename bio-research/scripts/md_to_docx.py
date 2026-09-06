#!/usr/bin/env python3
"""Convert a research report markdown file to DOCX."""

import argparse
import re
from pathlib import Path


def ensure_docx():
    try:
        import docx  # noqa: F401
    except ImportError:
        raise SystemExit(
            "python-docx is required. Install with: pip install --user python-docx"
        )


def parse_table(lines):
    """Parse a markdown table block into header + rows."""
    header = [c.strip() for c in lines[0].split("|") if c.strip()]
    rows = []
    for line in lines[2:]:
        cells = [c.strip() for c in line.split("|")]
        cells = [c for c in cells if c]
        if cells:
            rows.append(cells)
    return header, rows


def parse_markdown(text: str):
    blocks = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Headings
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            blocks.append(("heading", level, m.group(2).strip()))
            i += 1
            continue

        # Horizontal rule / page break
        if re.match(r"^\s*[-*]{3,}\s*$", stripped):
            blocks.append(("rule",))
            i += 1
            continue

        # Table
        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            header, rows = parse_table(table_lines)
            blocks.append(("table", header, rows))
            continue

        # List item
        if re.match(r"^[-*]\s+", stripped):
            items = []
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].strip()))
                i += 1
            blocks.append(("list", items))
            continue

        # Paragraph
        para_lines = []
        while i < len(lines) and lines[i].strip():
            para_lines.append(lines[i].strip())
            i += 1
        if para_lines:
            blocks.append(("paragraph", " ".join(para_lines)))
        else:
            i += 1

    return blocks


def render_inline(runs_parent, text: str):
    """Add text to a paragraph or table cell with inline formatting."""
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[\d+\])", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            runs_parent.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            runs_parent.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = runs_parent.add_run(part[1:-1])
            r.font.name = "Courier New"
        elif re.match(r"^\[\d+\]$", part):
            runs_parent.add_run(part).font.superscript = True
        else:
            runs_parent.add_run(part)


def md_to_docx(input_path: Path, output_path: Path):
    ensure_docx()
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Default style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    blocks = parse_markdown(input_path.read_text(encoding="utf-8"))

    for block in blocks:
        if block[0] == "heading":
            _, level, text = block
            p = doc.add_heading(level=level)
            render_inline(p, text)
        elif block[0] == "paragraph":
            _, text = block
            p = doc.add_paragraph()
            render_inline(p, text)
        elif block[0] == "list":
            _, items = block
            for item in items:
                p = doc.add_paragraph(style="List Bullet")
                render_inline(p, item)
        elif block[0] == "table":
            _, header, rows = block
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            hdr_cells = table.rows[0].cells
            for j, h in enumerate(header):
                render_inline(hdr_cells[j].paragraphs[0], h)
            for row in rows:
                row_cells = table.add_row().cells
                for j, cell in enumerate(row):
                    if j < len(row_cells):
                        render_inline(row_cells[j].paragraphs[0], cell)
        elif block[0] == "rule":
            doc.add_paragraph("_" * 50)

    # Center title if first paragraph is heading level 1
    if doc.paragraphs and doc.paragraphs[0].style.name.startswith("Heading 1"):
        doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"DOCX written to {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Convert Markdown report to DOCX")
    parser.add_argument("input", type=Path, help="Input Markdown file")
    parser.add_argument("output", type=Path, help="Output DOCX file")
    args = parser.parse_args()
    md_to_docx(args.input, args.output)


if __name__ == "__main__":
    main()
