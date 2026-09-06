#!/usr/bin/env python3
"""
Extract PPTX text and run structure as JSON.

Usage: python extract_pptx.py input.pptx [output.json]

Outputs JSON:
{
  "file_type": "pptx",
  "elements": [
    {
      "key": "(0, 3)",
      "type": "shape",
      "slide": 0,
      "shape": 3,
      "runs": 2,
      "text": "full text",
      "parts": ["run0", "run1"]
    },
    {
      "key": "(0, 5, 1, 2)",
      "type": "table_cell",
      "slide": 0,
      "shape": 5,
      "table_row": 1,
      "table_col": 2,
      "runs": 1,
      "text": "cell text",
      "parts": ["run0"]
    }
  ]
}
"""
import sys
import json
from pathlib import Path
from pptx import Presentation


def extract_pptx(pptx_path, out_path=None):
    prs = Presentation(pptx_path)
    elements = []

    for si, slide in enumerate(prs.slides):
        for sx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                tf = shape.text_frame

                total_runs = 0
                all_runs = []
                full_text_parts = []

                for pi, para in enumerate(tf.paragraphs):
                    para_runs = [r for r in para.runs if r.text.strip()]
                    if para_runs:
                        total_runs += len(para_runs)
                        all_runs.extend(para_runs)
                        para_text = ''.join(r.text for r in para_runs)
                        full_text_parts.append(para_text)

                full_text = '\n'.join(full_text_parts)

                if full_text.strip():
                    elements.append({
                        "key": f"({si}, {sx})",
                        "type": "shape",
                        "slide": si,
                        "shape": sx,
                        "runs": total_runs,
                        "text": full_text,
                        "parts": [r.text for r in all_runs]
                    })

            elif shape.has_table:
                for ri in range(len(shape.table.rows)):
                    for ci in range(len(shape.table.rows[ri].cells)):
                        cell = shape.table.cell(ri, ci)
                        cell_text = cell.text_frame.text.strip()
                        if cell_text:
                            cell_runs = [r for para in cell.text_frame.paragraphs
                                         for r in para.runs if r.text.strip()]
                            elements.append({
                                "key": f"({si}, {sx}, {ri}, {ci})",
                                "type": "table_cell",
                                "slide": si,
                                "shape": sx,
                                "table_row": ri,
                                "table_col": ci,
                                "runs": len(cell_runs),
                                "text": cell_text,
                                "parts": [r.text for r in cell_runs]
                            })

    result = {
        "file_type": "pptx",
        "source_file": str(pptx_path),
        "elements": elements,
        "total_elements": len(elements),
        "total_runs": sum(e["runs"] for e in elements)
    }

    json_str = json.dumps(result, ensure_ascii=False, indent=2)

    if out_path:
        Path(out_path).write_text(json_str, encoding='utf-8')
        print(f'Written to {out_path}')
    else:
        print(json_str)

    return result


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(f'Usage: python {sys.argv[0]} input.pptx [output.json]')
        sys.exit(1)

    pptx_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else None
    extract_pptx(pptx_path, out_path)
