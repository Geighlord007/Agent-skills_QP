# 合一写回器 v3：DOCX（WIR / surgical 环境自适应）与 PPTX（含组合形状、表格、备注）
# 用法: python write_v3.py input.docx|input.pptx output translated.json
# DOCX 引擎按 engine_select.py 的判定分发：WIR（Linux）或 surgical（其余平台）；
# PPTX 按 key 中的形状路径（含组合形状子级）定位元素，仅写 run 文本，排版一律不碰。
import json
import re
import subprocess
import sys
from pathlib import Path

from pptx import Presentation

SCRIPTS = Path(__file__).resolve().parent


def resolve_shape(shapes, path):
    sh = shapes[path[0]]
    for idx in path[1:]:
        sh = sh.shapes[idx]
    return sh


def set_runs_in_text_frame(text_frame, parts):
    run_map = [(pi, ri) for pi, para in enumerate(text_frame.paragraphs)
               for ri, _ in enumerate(para.runs)]
    if len(run_map) != len(parts):
        raise ValueError("run 数量 %d 与 parts 数量 %d 不一致" % (len(run_map), len(parts)))
    for pi, ri in run_map:
        text_frame.paragraphs[pi].runs[ri].text = ""
    for i, (pi, ri) in enumerate(run_map):
        text_frame.paragraphs[pi].runs[ri].text = parts[i]


def write_pptx(input_path, output_path, data):
    prs = Presentation(input_path)
    written = 0
    for elem in data["elements"]:
        key, parts = elem["key"], elem.get("parts", [])
        if key.startswith("s_"):
            segs = key.split("_")[1:]
            sh = resolve_shape(prs.slides[int(segs[0])].shapes, [int(x) for x in segs[1:]])
            if not sh.has_text_frame:
                raise ValueError("%s：形状没有文本框" % key)
            set_runs_in_text_frame(sh.text_frame, parts)
        elif key.startswith("t_"):
            segs = key.split("_")[1:]
            slide_idx = int(segs[0])
            row, col = int(segs[-2]), int(segs[-1])
            sh = resolve_shape(prs.slides[slide_idx].shapes, [int(x) for x in segs[1:-2]])
            set_runs_in_text_frame(sh.table.cell(row, col).text_frame, parts)
        elif key.startswith("n_"):
            _, slide_idx, para_idx = key.split("_")
            para = prs.slides[int(slide_idx)].notes_slide.notes_text_frame.paragraphs[int(para_idx)]
            if len(para.runs) != len(parts):
                raise ValueError("%s：run 数量与 parts 数量不一致" % key)
            for r, p in zip(para.runs, parts):
                r.text = ""
            for r, p in zip(para.runs, parts):
                r.text = p
        else:
            raise ValueError("无法识别的 key: %s" % key)
        written += 1
    prs.save(output_path)
    print("PPTX 写回 %d 个元素 → %s" % (written, output_path))


def write_docx(input_path, output_path, json_path):
    decision = subprocess.run(
        [sys.executable, str(SCRIPTS / "engine_select.py")],
        capture_output=True, text=True, check=True)
    engine = json.loads(decision.stdout)["recommended"]
    writer = SCRIPTS / ("write_docx_wir.py" if engine == "wir" else "write_docx_surgical.py")
    print("DOCX 写回引擎: %s（%s）" % (engine, writer.name))
    subprocess.run([sys.executable, str(writer), str(input_path), str(output_path), str(json_path)],
                   check=True)


def main():
    if len(sys.argv) < 4:
        raise SystemExit("用法: python write_v3.py input.docx|input.pptx output translated.json")
    src, dst, js = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
    data = json.loads(js.read_text(encoding="utf-8"))
    if src.suffix.lower() == ".pptx":
        write_pptx(src, dst, data)
    elif src.suffix.lower() == ".docx":
        write_docx(src, dst, js)
    else:
        raise SystemExit("不支持的文件类型: %s（仅 .docx / .pptx）" % src.suffix)


if __name__ == "__main__":
    main()
