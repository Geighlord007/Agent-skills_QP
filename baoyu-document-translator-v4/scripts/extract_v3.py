# 合一提取器 v3：DOCX（全部 story）与 PPTX（组合形状、表格、备注）→ keyed JSON
# 用法: python extract_v3.py input.docx|input.pptx extracted.json [--no-notes]
# PPTX 键名: s_{slide}_{形状路径（下划线连接，含组合形状子级）}, t_{slide}_{路径}_{行}_{列}, n_{slide}_{段号}
# DOCX 键名与 extract_docx_v2 完全一致（写回器按 key 定位元素）
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE


def run_style(run):
    # run 格式签名：粗体/斜体/下划线/字号/颜色，供质检比对格式是否保持
    f = run.font
    sig = []
    if f.bold:
        sig.append("b")
    if f.italic:
        sig.append("i")
    if f.underline:
        sig.append("u")
    if f.size is not None:
        sig.append("sz%d" % f.size.pt)
    color = f.color
    if color is not None and color.type == MSO_COLOR_TYPE.RGB:
        sig.append("c{}".format(str(color.rgb)))
    return ",".join(sig) or "-"


def collect_text_frame(tf):
    # 返回 (run 文本列表, 格式签名列表, 分段累计 run 数, 全文)；空段（无 run）不计入
    runs, styles, para_splits, texts = [], [], [], []
    acc = 0
    for para in tf.paragraphs:
        para_runs = [r for r in para.runs]
        if not para_runs:
            continue
        acc += len(para_runs)
        para_splits.append(acc)
        runs.extend(r.text for r in para_runs)
        styles.extend(run_style(r) for r in para_runs)
        texts.append("".join(r.text for r in para_runs))
    return runs, styles, para_splits, "\n".join(texts)


def make_element(key, elem_type, story, runs, styles, para_splits, text, meta):
    e = {"key": key, "type": elem_type, "story": story}
    e.update(meta)
    e["runs"] = len(runs)
    e["text"] = text
    e["parts"] = runs
    if any(s != "-" for s in styles):
        e["run_styles"] = styles
    if len(para_splits) > 1:
        e["para_splits"] = para_splits
    return e


def pptx_walk(shapes, si, path, elements):
    for idx, sh in enumerate(shapes):
        cur = path + [idx]
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            pptx_walk(sh.shapes, si, cur, elements)
            continue
        if sh.has_table:
            for ri, row in enumerate(sh.table.rows):
                for ci, cell in enumerate(row.cells):
                    runs, styles, splits, text = collect_text_frame(cell.text_frame)
                    if not text.strip():
                        continue
                    key = "t_%d_%s_%d_%d" % (si, "_".join(map(str, cur)), ri, ci)
                    meta = {"slide": si, "shape_path": cur, "table_row": ri, "table_col": ci}
                    elements.append(make_element(key, "table_cell", "slide",
                                                 runs, styles, splits, text, meta))
            continue
        if sh.has_text_frame:
            runs, styles, splits, text = collect_text_frame(sh.text_frame)
            if not text.strip():
                continue
            key = "s_%d_%s" % (si, "_".join(map(str, cur)))
            meta = {"slide": si, "shape_path": cur}
            elements.append(make_element(key, "shape", "slide",
                                         runs, styles, splits, text, meta))


def extract_pptx(pptx_path, include_notes=True):
    prs = Presentation(pptx_path)
    elements = []
    for si, slide in enumerate(prs.slides):
        pptx_walk(slide.shapes, si, [], elements)
        if include_notes and slide.has_notes_slide:
            ntf = slide.notes_slide.notes_text_frame
            for pi, para in enumerate(ntf.paragraphs):
                para_runs = [r for r in para.runs]
                if not para_runs:
                    continue
                text = "".join(r.text for r in para_runs)
                if not text.strip():
                    continue
                key = "n_%d_%d" % (si, pi)
                meta = {"slide": si, "note_para": pi}
                elements.append(make_element(key, "paragraph", "notes",
                                             [r.text for r in para_runs],
                                             [run_style(r) for r in para_runs],
                                             [len(para_runs)], text, meta))
    return {
        "file_type": "pptx",
        "source_file": str(pptx_path),
        "elements": elements,
        "total_elements": len(elements),
        "total_runs": sum(e["runs"] for e in elements),
    }


def extract_docx_part(docx_path, out_path):
    # DOCX 侧沿用 v2 提取逻辑（全部 story、域 run 守护、超链接），键名保持不变
    from extract_docx_v2 import extract_docx
    return extract_docx(docx_path, str(out_path))


def main():
    if len(sys.argv) < 3:
        raise SystemExit("用法: python extract_v3.py input.docx|input.pptx extracted.json [--no-notes]")
    src = Path(sys.argv[1])
    out = Path(sys.argv[2])
    suffix = src.suffix.lower()
    if suffix == ".docx":
        extract_docx_part(src, out)
        return
    if suffix != ".pptx":
        raise SystemExit("不支持的文件类型: %s（仅 .docx / .pptx）" % suffix)
    include_notes = "--no-notes" not in sys.argv
    data = extract_pptx(src, include_notes)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Written %s: %d elements, %d runs" % (out, data["total_elements"], data["total_runs"]))


if __name__ == "__main__":
    main()
