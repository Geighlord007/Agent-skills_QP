#!/usr/bin/env python3
"""
Extract PPTX text and run structure as keyed JSON, including shapes,
table cells, and optional notes.

Usage: python extract_pptx_v2.py input.pptx [output.json] [--include-notes]
"""
import sys
import json
from pathlib import Path
from pptx import Presentation


def extract_pptx(pptx_path, out_path=None, include_notes=False):
    prs = Presentation(pptx_path)
    elements = []

    for si, slide in enumerate(prs.slides):
        for sx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                tf = shape.text_frame

                all_runs = []
                full_text_parts = []

                for para in tf.paragraphs:
                    para_runs = [r for r in para.runs]
                    if para_runs:
                        all_runs.extend(para_runs)
                        para_text = "".join(r.text for r in para_runs)
                        full_text_parts.append(para_text)

                full_text = "\n".join(full_text_parts)

                if full_text.strip():
                    elements.append({
                        "key": f"s_{si}_{sx}",
                        "type": "shape",
                        "story": "slide",
                        "slide": si,
                        "shape": sx,
                        "runs": len(all_runs),
                        "text": full_text,
                        "parts": [r.text for r in all_runs],
                    })

            elif shape.has_table:
                for ri, row in enumerate(shape.table.rows):
                    for ci, cell in enumerate(row.cells):
                        cell_tf = cell.text_frame
                        cell_runs = [r for para in cell_tf.paragraphs for r in para.runs]
                        cell_text = cell_tf.text.strip()
                        if cell_text:
                            elements.append({
                                "key": f"t_{si}_{sx}_{ri}_{ci}",
                                "type": "table_cell",
                                "story": "slide",
                                "slide": si,
                                "shape": sx,
                                "table_row": ri,
                                "table_col": ci,
                                "runs": len(cell_runs),
                                "text": cell_tf.text,
                                "parts": [r.text for r in cell_runs],
                            })

        if include_notes and slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            for pi, para in enumerate(notes.paragraphs):
                para_runs = [r for r in para.runs]
                para_text = "".join(r.text for r in para_runs).strip()
                if para_text:
                    elements.append({
                        "key": f"n_{si}_{pi}",
                        "type": "paragraph",
                        "story": "notes",
                        "slide": si,
                        "index": pi,
                        "runs": len(para_runs),
                        "text": para_text,
                        "parts": [r.text for r in para_runs],
                    })

    result = {
        "file_type": "pptx",
        "source_file": str(pptx_path),
        "elements": elements,
        "total_elements": len(elements),
        "total_runs": sum(e["runs"] for e in elements),
    }

    json_str = json.dumps(result, ensure_ascii=False, indent=2)

    if out_path:
        Path(out_path).write_text(json_str, encoding="utf-8")
        print(f"Written to {out_path}")
    else:
        print(json_str)

    return result


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(f"Usage: python {sys.argv[0]} input.pptx [output.json] [--include-notes]")
        sys.exit(1)

    pptx_path = sys.argv[1]
    out_path = None
    include_notes = False

    for arg in sys.argv[2:]:
        if arg == "--include-notes":
            include_notes = True
        elif not out_path:
            out_path = arg

    extract_pptx(pptx_path, out_path, include_notes)
