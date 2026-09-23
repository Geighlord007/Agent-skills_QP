# -*- coding: utf-8 -*-
"""Build a structural fixture docx for selftest: PAGE field + ptab (footer), first-page
header, horizontal+vertical merged cells, tracked-insertion run, multi-run paragraph.
Usage: python build_fixture.py out.docx"""
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
doc.add_heading("Fixture Report", level=1)
p = doc.add_paragraph()
r1 = p.add_run("Bold lead: ")
r1.bold = True
p.add_run("normal continuation with numbers 42 and 4.5%.")

# tracked insertion run
p2 = doc.add_paragraph("Base sentence. ")
ins = OxmlElement("w:ins")
ins.set(qn("w:id"), "77")
ins.set(qn("w:author"), "selftest")
ins.set(qn("w:date"), "2026-09-22T00:00:00Z")
rins = OxmlElement("w:r")
t = OxmlElement("w:t")
t.text = "Inserted clause stays translatable."
rins.append(t)
ins.append(rins)
p2._p.append(ins)

# table with horizontal merge (row0: c0+c1 merged) and vertical merge (col2 rows1-2)
tbl = doc.add_table(rows=3, cols=3)
tbl.cell(0, 0).merge(tbl.cell(0, 1))
tbl.cell(0, 0).text = "Merged header span"
tbl.cell(0, 2).text = "H2"
tbl.cell(1, 0).text = "a"
tbl.cell(1, 1).text = "b"
tbl.cell(1, 2).text = "vertical span start"
tbl.cell(2, 0).text = "c"
tbl.cell(2, 1).text = "d"
tbl.cell(1, 2).merge(tbl.cell(2, 2))
doc.add_paragraph("Trailing paragraph after table.")

# footer: ptab + PAGE field
sec = doc.sections[0]
fp = sec.footer.paragraphs[0]
fr = fp.add_run("Footer left")
rt = fp.add_run()
rt._r.append(OxmlElement("w:ptab"))
fb = OxmlElement("w:r")
fc1 = OxmlElement("w:fldChar"); fc1.set(qn("w:fldCharType"), "begin"); fb.append(fc1)
fi = OxmlElement("w:r")
it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
fi.append(it)
fs = OxmlElement("w:r")
fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "separate"); fs.append(fc2)
fr2 = OxmlElement("w:r")
t2 = OxmlElement("w:t"); t2.text = "1"; fr2.append(t2)
fe = OxmlElement("w:r")
fc3 = OxmlElement("w:fldChar"); fc3.set(qn("w:fldCharType"), "end"); fe.append(fc3)
for el in (fb, fi, fs, fr2, fe):
    fp._p.append(el)

# distinct first-page header
sec.different_first_page_header_footer = True
sec.first_page_header.paragraphs[0].add_run("COVER HEADER LINE")

# body text box (pict > shape > txbxContent)
p_tb = doc.add_paragraph()
run_tb = p_tb.add_run()
pict = OxmlElement("w:pict")
shape = OxmlElement("w:shape")
tb = OxmlElement("w:txbxContent")
tp = OxmlElement("w:p")
tr_ = OxmlElement("w:r")
tt = OxmlElement("w:t")
tt.text = "Textbox line one."
tr_.append(tt)
tp.append(tr_)
tb.append(tp)
shape.append(tb)
pict.append(shape)
run_tb._r.append(pict)

# SDT content control paragraph
sdt = OxmlElement("w:sdt")
sdtc = OxmlElement("w:sdtContent")
sp = OxmlElement("w:p")
sr = OxmlElement("w:r")
st = OxmlElement("w:t")
st.text = "SDT caption line."
sr.append(st)
sp.append(sr)
sdtc.append(sp)
sdt.append(sdtc)
doc.element.body.append(sdt)

# run carrying MULTIPLE w:t nodes (WPS artifact)
pm = doc.add_paragraph()
rm = pm.add_run()
t1 = OxmlElement("w:t")
t1.text = "First half "
rm._r.append(t1)
t2 = OxmlElement("w:t")
t2.text = "second half in same run."
rm._r.append(t2)

# paragraph whose translatable text lives in a hyperlink run
ph = doc.add_paragraph("Prefix ")
hl = OxmlElement("w:hyperlink")
hl.set(qn("w:anchor"), "_selftest")
hr = OxmlElement("w:r")
ht = OxmlElement("w:t")
ht.text = "hyperlinked phrase"
hr.append(ht)
hl.append(hr)
ph._p.append(hl)

# two-paragraph cell (para_splits)
tbl2 = doc.add_table(rows=1, cols=1)
c2 = tbl2.cell(0, 0)
c2.text = "Line one of cell"
c2.add_paragraph("Line two of cell")

# cover-style paragraph: 4 runs, 4 distinct rPr, short (consolidation target)
from docx.shared import Pt
pc = doc.add_paragraph()
for i, txt in enumerate(["Cover ", "Title ", "Block ", "Here"]):
    r = pc.add_run(txt)
    r.bold = (i % 2 == 0)
    r.italic = (i % 2 == 1)
    r.font.size = Pt(12 + i)

# textbox paragraph with [Alpha][tab][Beta] runs (W1 alignment regression)
p_tb2 = doc.add_paragraph()
r_tb2 = p_tb2.add_run()
pict2 = OxmlElement("w:pict")
shape2 = OxmlElement("w:shape")
tb2 = OxmlElement("w:txbxContent")
tp2 = OxmlElement("w:p")
ra = OxmlElement("w:r")
ta = OxmlElement("w:t")
ta.text = "Alpha sufficiently long textbox sentence one"
ra.append(ta)
rtab = OxmlElement("w:r")
rtab.append(OxmlElement("w:tab"))
rb = OxmlElement("w:r")
tb_ = OxmlElement("w:t")
tb_.text = "Beta sufficiently long textbox sentence two"
rb.append(tb_)
for el in (ra, rtab, rb):
    tp2.append(el)
tb2.append(tp2)
shape2.append(tb2)
pict2.append(shape2)
r_tb2._r.append(pict2)

# body caption with SEQ field (field-result run must be excluded from parts)
pcap = doc.add_paragraph()
pcap.add_run("Table ")
fb = OxmlElement("w:r")
fc = OxmlElement("w:fldChar")
fc.set(qn("w:fldCharType"), "begin")
fb.append(fc)
fi = OxmlElement("w:r")
it = OxmlElement("w:instrText")
it.text = " SEQ Table \\* ARABIC "
fi.append(it)
fs = OxmlElement("w:r")
fc2 = OxmlElement("w:fldChar")
fc2.set(qn("w:fldCharType"), "separate")
fs.append(fc2)
fr = OxmlElement("w:r")
ft = OxmlElement("w:t")
ft.text = "7"
fr.append(ft)
fe = OxmlElement("w:r")
fc3 = OxmlElement("w:fldChar")
fc3.set(qn("w:fldCharType"), "end")
fe.append(fc3)
for el in (fb, fi, fs, fr, fe):
    pcap._p.append(el)
pcap.add_run(". Caption with field number.")

doc.save(sys.argv[1])
print("fixture written:", sys.argv[1])
