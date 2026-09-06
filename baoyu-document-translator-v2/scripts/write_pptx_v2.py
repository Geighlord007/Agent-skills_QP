#!/usr/bin/env python3
"""
Write keyed translations from JSON back to PPTX, including shapes and table cells.

Usage: python write_pptx_v2.py input.pptx output.pptx translations.json
"""
import sys
import json
import re
from pathlib import Path
from pptx import Presentation


def get_run_map(text_frame):
    """Return list of (para_idx, run_idx) for all runs in a text frame."""
    run_map = []
    for pi, para in enumerate(text_frame.paragraphs):
        for ri, run in enumerate(para.runs):
            run_map.append((pi, ri))
    return run_map


def set_runs_in_text_frame(text_frame, parts):
    """Clear and rewrite runs in a text frame, preserving formatting."""
    run_map = get_run_map(text_frame)

    # Clear
    for pi, ri in run_map:
        text_frame.paragraphs[pi].runs[ri].text = ""

    # Write back
    for i, (pi, ri) in enumerate(run_map):
        if i < len(parts):
            text_frame.paragraphs[pi].runs[ri].text = parts[i]


def write_shape(shape, parts):
    if shape.has_text_frame:
        set_runs_in_text_frame(shape.text_frame, parts)
        return True
    return False


def write_table_cell(table, row, col, parts):
    cell = table.cell(row, col)
    set_runs_in_text_frame(cell.text_frame, parts)
    return True


def write_pptx(input_path, output_path, json_path):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    translations = {e["key"]: e.get("parts", []) for e in data.get("elements", [])}
    prs = Presentation(input_path)

    written = 0
    skipped = 0
    errors = []

    for elem in data.get("elements", []):
        key = elem.get("key", "")
        parts = translations.get(key, [])

        # Shape key: s_slide_shape
        m = re.match(r"^s_(\d+)_(\d+)$", key)
        if m:
            slide_idx, shape_idx = int(m.group(1)), int(m.group(2))
            try:
                shape = prs.slides[slide_idx].shapes[shape_idx]
                if write_shape(shape, parts):
                    written += 1
                else:
                    skipped += 1
            except Exception as e:
                errors.append(f"{key}: {e}")
                skipped += 1
            continue

        # Table cell key: t_slide_shape_row_col
        m = re.match(r"^t_(\d+)_(\d+)_(\d+)_(\d+)$", key)
        if m:
            slide_idx, shape_idx, row, col = map(int, m.groups())
            try:
                shape = prs.slides[slide_idx].shapes[shape_idx]
                if shape.has_table:
                    write_table_cell(shape.table, row, col, parts)
                    written += 1
                else:
                    skipped += 1
            except Exception as e:
                errors.append(f"{key}: {e}")
                skipped += 1
            continue

        # Notes key: n_slide_pIndex (optional)
        m = re.match(r"^n_(\d+)_(\d+)$", key)
        if m:
            slide_idx, para_idx = int(m.group(1)), int(m.group(2))
            try:
                notes = prs.slides[slide_idx].notes_text_frame
                if notes:
                    para = notes.paragraphs[para_idx]
                    set_runs_in_text_frame(para, parts)
                    written += 1
                else:
                    skipped += 1
            except Exception as e:
                errors.append(f"{key}: {e}")
                skipped += 1
            continue

        skipped += 1

    prs.save(output_path)

    print(f"\n{'='*50}")
    print("PPTX Write Complete (v2)")
    print(f"{'='*50}")
    print(f"  Written: {written} elements")
    if skipped > 0:
        print(f"  Skipped: {skipped} elements")
    if errors:
        print(f"  Errors: {len(errors)}")
        for e in errors[:10]:
            print(f"    - {e}")
    print(f"  Output: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("Usage: python write_pptx_v2.py <input.pptx> <output.pptx> <translations.json>")
        sys.exit(1)

    write_pptx(sys.argv[1], sys.argv[2], sys.argv[3])
