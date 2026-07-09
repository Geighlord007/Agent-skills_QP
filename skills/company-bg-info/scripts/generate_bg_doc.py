#!/usr/bin/env python3
"""
Generate a BG-info Word document from structured JSON data.
Usage: python generate_bg_doc.py --input data.json --output /path/to/Company_BG_info.docx

The JSON schema is documented in references/json-schema.md.
"""

import argparse
import json
import os
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError as e:
    print("ERROR: python-docx is required. Install: pip install --user python-docx")
    sys.exit(1)


# ── Colour palette ───────────────────────────────────────────────────────────
DARK_GREY = RGBColor(0x33, 0x33, 0x33)
GREEN = RGBColor(0x61, 0x7D, 0x24)
LINK_BLUE = RGBColor(0x05, 0x63, 0xC1)
LIGHT_GREY = "E5E5E5"
MED_GREY = RGBColor(0x66, 0x66, 0x66)

FONT_NAME = "Schibsted Grotesk"


def add_hyperlink(paragraph, text, url):
    """Insert a blue, underlined hyperlink into a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    hyperlink.set(qn("w:history"), "1")
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def set_para_font(p, name=FONT_NAME, size=Pt(10), color=DARK_GREY, bold=False, italic=False):
    for run in p.runs:
        run.font.name = name
        run.font.size = size
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def shade_cell(cell, fill=LIGHT_GREY):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_para_font(p, bold=True, size=Pt(9))
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(hdr[i])
    for row_data in rows:
        row = table.add_row().cells
        for i, txt in enumerate(row_data):
            row[i].text = str(txt)
            for p in row[i].paragraphs:
                set_para_font(p, size=Pt(9))
            row[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def add_styled(doc, text, style="Normal", bold=False, size=Pt(10),
               color=DARK_GREY, align=WD_ALIGN_PARAGRAPH.LEFT,
               space_after=Pt(6), italic=False):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_after = space_after
    p.paragraph_format.space_before = Pt(0)
    if text:
        p.add_run(text)
    set_para_font(p, size=size, color=color, bold=bold, italic=italic)
    return p


def add_image_with_caption(doc, image_path, caption, width=Inches(5.5)):
    if not os.path.exists(image_path):
        add_styled(doc, f"[Image not found: {image_path}]", italic=True, color=MED_GREY, size=Pt(8))
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_img.add_run()
    run.add_picture(image_path, width=width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(caption)
    set_para_font(cap, size=Pt(8), italic=True, color=MED_GREY)


def build_document(data, output_path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)

    # Normal style
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10)
    normal.font.color.rgb = DARK_GREY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    # Heading styles
    for name in ["Heading 1", "Heading 2"]:
        s = doc.styles[name]
        s.font.name = FONT_NAME
        s.font.size = Pt(10)
        s.font.bold = True
        s.font.color.rgb = GREEN
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(6)
        s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    # ── 1. Title ────────────────────────────────────────────────────────────
    add_styled(doc, f"{data['company_name']} BG info", size=Pt(14), bold=True, space_after=Pt(2))

    # ── 2. Main heading ───────────────────────────────────────────────────────
    add_styled(doc, data["company_name"], size=Pt(18), bold=True, space_after=Pt(8))

    # ── 3. Company overview ─────────────────────────────────────────────────
    add_styled(doc, data.get("overview", ""))

    # ── 4. Strategic context ──────────────────────────────────────────────────
    add_styled(doc, "Strategic context", size=Pt(10), bold=True, color=GREEN, space_after=Pt(6))
    add_styled(doc, data.get("strategic_context", ""))
    if data.get("strategic_image"):
        add_image_with_caption(doc, data["strategic_image"], data.get("strategic_image_caption", "Figure 1."))

    # ── 5. Operating divisions ──────────────────────────────────────────────
    add_styled(doc, "Operating divisions", size=Pt(10), bold=True, color=GREEN, space_after=Pt(6))
    divs = data.get("divisions", [])
    if divs:
        add_table(doc,
            ["Division", "Description", "Key brands / products", "Revenue / EBITDA"],
            [[d.get("name",""), d.get("description",""), d.get("brands",""), d.get("revenue","")] for d in divs])
    else:
        add_styled(doc, "No division data provided.", italic=True, color=MED_GREY)

    # ── 6. Selected brands ──────────────────────────────────────────────────
    brands = data.get("brands", {})
    if brands:
        add_styled(doc, "Selected ingredient / product brands", size=Pt(10), bold=True, color=GREEN, space_after=Pt(6))
        for category, items in brands.items():
            add_styled(doc, category, style="Heading 2", bold=True, color=GREEN, space_after=Pt(4))
            for item in items:
                add_styled(doc, f"• {item}")

    # ── 7. Key financials ───────────────────────────────────────────────────
    fins = data.get("financials", [])
    if fins:
        add_styled(doc, "Key financials", size=Pt(10), bold=True, color=GREEN, space_after=Pt(6))
        headers = ["Metric"] + [f["year"] for f in fins]
        rows = []
        for metric in fins[0].get("metrics", {}).keys():
            row = [metric]
            for f in fins:
                row.append(f.get("metrics", {}).get(metric, ""))
            rows.append(row)
        add_table(doc, headers, rows)

    # ── 8. Alternative proteins / precision fermentation ────────────────────
    add_styled(doc, "Alternative proteins / precision fermentation / relevant technology layout",
               size=Pt(10), bold=True, color=GREEN, space_after=Pt(6))
    alt = data.get("alternative_proteins", {})
    if not alt:
        add_styled(doc, "No public internal platform announced; strategic interest not yet signalled.")
    else:
        for key in ["public_statements", "initiatives", "partnerships", "investments", "ma_deals"]:
            val = alt.get(key)
            if val:
                add_styled(doc, key.replace("_", " ").title(), style="Heading 2", bold=True, color=GREEN, space_after=Pt(4))
                if isinstance(val, list):
                    for v in val:
                        add_styled(doc, f"• {v}")
                else:
                    add_styled(doc, str(val))

    # ── 9. Recent M&A / transactions ──────────────────────────────────────
    ma = data.get("ma_transactions", [])
    if ma:
        add_styled(doc, "Recent M&A / transactions", size=Pt(10), bold=True, color=GREEN, space_after=Pt(6))
        add_table(doc,
            ["Year", "Target / asset", "Transaction amount", "Description"],
            [[m.get("year",""), m.get("target",""), m.get("amount",""), m.get("description","")] for m in ma])

    # ── 10. Innovation & partnership landscape ──────────────────────────────
    add_styled(doc, "Innovation & partnership landscape", size=Pt(10), bold=True, color=GREEN, space_after=Pt(6))
    for key in ["open_innovation", "relevant_products", "pf_stance", "collaboration_hook"]:
        val = data.get("innovation", {}).get(key)
        if val:
            add_styled(doc, key.replace("_", " ").title(), style="Heading 2", bold=True, color=GREEN, space_after=Pt(4))
            add_styled(doc, str(val))

    # ── 11. Sources ─────────────────────────────────────────────────────────
    add_styled(doc, "Sources", size=Pt(10), bold=True, color=GREEN, space_after=Pt(6))
    for src in data.get("sources", []):
        p = doc.add_paragraph()
        add_hyperlink(p, src.get("label", src["url"]), src["url"])
        set_para_font(p)

    doc.save(output_path)
    print(f"Saved: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Generate BG-info DOCX from JSON")
    parser.add_argument("--input", "-i", required=True, help="Path to JSON data file")
    parser.add_argument("--output", "-o", required=True, help="Output DOCX path")
    args = parser.parse_args()

    with open(args.input, "r", encoding="utf-8") as f:
        data = json.load(f)

    build_document(data, args.output)


if __name__ == "__main__":
    main()
