#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成会议纪要 docx - 带图片嵌入和格式化"""
import os, sys, re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image

def normalize_image(path):
    """用 PIL 重新保存图片，避免 python-docx 的 JPEG 嵌入 bug"""
    try:
        img = Image.open(path)
        img = img.convert('RGB')  # 关键！避免 RGBA 问题
        img.save(path, 'JPEG', quality=85)
    except Exception as e:
        print(f"    [warn] 图片 normalize 失败: {path} - {e}")

def add_company_section(doc, company_name, speaker, summary, key_points, image_folder, image_files):
    """添加一家公司的章节"""
    # 公司名 - 加粗 Heading
    h = doc.add_heading(level=1)
    h.clear()
    run = h.add_run(company_name)
    run.bold = True
    run.font.size = Pt(14)
    
    # 演讲人
    if speaker:
        p = doc.add_paragraph()
        p.add_run(f"演讲人：{speaker}").italic = True
    
    # 核心亮点 - 加粗
    if key_points:
        for point in key_points:
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(point)
            run.bold = True
    
    # 详细总结
    if summary:
        doc.add_paragraph(summary)
    
    # 图片
    if image_files:
        for img_file in image_files:
            img_path = os.path.join(image_folder, img_file)
            if os.path.exists(img_path):
                # 先 normalize 避免 docx 嵌入 bug
                normalize_image(img_path)
                try:
                    doc.add_picture(img_path, width=Inches(5.5))
                    last_para = doc.paragraphs[-1]
                    last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception as e:
                    doc.add_paragraph(f"[图片无法嵌入: {img_file}]")
    
    doc.add_paragraph()  # 空行分隔

def create_docx(company_data, output_path):
    """
    company_data: list of dicts:
        {
            'name': '公司名',
            'speaker': '演讲人',
            'summary': '详细总结（可为空）',
            'key_points': ['亮点1', '亮点2'],
            'image_folder': '图片文件夹路径',
            'image_files': ['img1.jpg', 'img2.jpg']
        }
    """
    doc = Document()
    
    # 标题
    title = doc.add_heading('会议纪要', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    for company in company_data:
        add_company_section(
            doc,
            company.get('name', ''),
            company.get('speaker', ''),
            company.get('summary', ''),
            company.get('key_points', []),
            company.get('image_folder', ''),
            company.get('image_files', [])
        )
    
    doc.save(output_path)
    print(f"文档已保存: {output_path}")
    return output_path

if __name__ == '__main__':
    # 示例用法
    example_data = [
        {
            'name': '恩和科技',
            'speaker': '宋任剑，AI & Computation总监',
            'summary': '恩和科技是一家专注于AI与合成生物学结合的技术平台公司...',
            'key_points': ['核心亮点1', '核心亮点2'],
            'image_folder': r'D:\WPS Software\工作文件夹\Claw\会议记录_2026-03-31\恩和科技',
            'image_files': ['IMG_001.jpg', 'IMG_002.jpg']
        }
    ]
    
    output = r'D:\WPS Software\工作文件夹\Claw\会议记录_2026-03-31\会议记录_2026-03-31_整理纪要.docx'
    create_docx(example_data, output)
