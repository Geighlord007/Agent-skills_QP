#!/usr/bin/env python3
"""
Write translations from validated JSON to PPTX.

Usage:
    python write_pptx.py input.pptx output.pptx translations.json

Strategy: Clear all runs first, then write translations back by run index.
Preserves font properties (size, color, bold, etc.) on existing runs.
"""
import sys
import json
import re
from pathlib import Path
from pptx import Presentation


def get_run_map(shape):
    """Return list of (para_idx, run_idx) tuples for all runs with text."""
    if not shape.has_text_frame:
        return []
    tf = shape.text_frame
    run_map = []
    for pi, para in enumerate(tf.paragraphs):
        for ri, run in enumerate(para.runs):
            if run.text.strip():
                run_map.append((pi, ri))
    return run_map


def clear_all_runs(shape):
    """Clear text from all runs while preserving paragraph structure."""
    if not shape.has_text_frame:
        return
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            run.text = ''


def set_text_run_by_run(shape, parts):
    """Write text parts to runs, preserving font properties."""
    if not shape.has_text_frame:
        return

    run_map = get_run_map(shape)
    clear_all_runs(shape)

    if len(parts) != len(run_map):
        print(f"  WARNING: parts({len(parts)}) != run_map({len(run_map)})")

    tf = shape.text_frame
    for i, (pi, ri) in enumerate(run_map):
        if i < len(parts):
            tf.paragraphs[pi].runs[ri].text = parts[i]
        else:
            print(f"  WARNING: no text for run {i} ({pi},{ri})")


def write_pptx(input_path, output_path, json_path):
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    translations = {e["key"]: e.get("parts", []) for e in data.get("elements", [])}

    prs = Presentation(input_path)
    written = 0
    skipped = 0

    for entry in data.get("elements", []):
        key = entry.get("key", "")
        slide_idx = entry.get("slide")
        shape_idx = entry.get("shape")
        parts = translations.get(key, [])

        if slide_idx is None or shape_idx is None:
            match = re.match(r"\((\d+),\s*(\d+)\)", key)
            if match:
                slide_idx, shape_idx = int(match.group(1)), int(match.group(2))

        if slide_idx is None or shape_idx is None:
            skipped += 1
            continue

        try:
            slide = prs.slides[slide_idx]
            shape = slide.shapes[shape_idx]

            if shape.has_text_frame:
                set_text_run_by_run(shape, parts)
                written += 1
            else:
                skipped += 1

        except (IndexError, Exception) as e:
            print(f"  ERROR {key}: {e}")
            skipped += 1

    prs.save(output_path)

    print(f"\n{'='*50}")
    print(f"PPTX Write Complete")
    print(f"{'='*50}")
    print(f"  Written: {written} shapes")
    if skipped > 0:
        print(f"  Skipped: {skipped} shapes")
    print(f"  Output: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("Usage: python write_pptx.py <input.pptx> <output.pptx> <translations.json>")
        sys.exit(1)

    write_pptx(sys.argv[1], sys.argv[2], sys.argv[3])
