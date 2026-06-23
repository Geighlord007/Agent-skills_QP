#!/usr/bin/env python3
"""
Write translations from validated JSON to DOCX.

Usage:
    python write_docx.py input.docx output.docx translations.json

Strategy: Clear all runs first, then write translations back by run index.
"""
import sys
import json
from pathlib import Path
from docx import Document


def clear_all_runs(para):
    """Clear text from all runs in a paragraph."""
    for run in para.runs:
        run.text = ''


def write_to_paragraph(para, parts):
    """Write parts to paragraph runs, preserving format."""
    runs = para.runs
    if not runs:
        if parts:
            para.add_run(parts[0])
        return

    # Clear first
    clear_all_runs(para)

    # Write back
    for i, run in enumerate(runs):
        if i < len(parts):
            run.text = parts[i]


def write_to_cell(cell, parts):
    """Write parts to cell runs across all paragraphs."""
    all_runs = [r for para in cell.paragraphs for r in para.runs]
    if not all_runs:
        if parts and cell.paragraphs:
            cell.paragraphs[0].add_run(parts[0])
        return

    # Clear first
    for run in all_runs:
        run.text = ''

    # Write back
    for i, run in enumerate(all_runs):
        if i < len(parts):
            run.text = parts[i]


def write_docx(input_path, output_path, json_path):
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    translations = {e["key"]: e.get("parts", []) for e in data.get("elements", [])}

    doc = Document(input_path)
    para_map = {id(p._p): p for p in doc.paragraphs}
    table_map = {id(t._tbl): t for t in doc.tables}

    para_idx = 0
    table_idx = 0
    written = 0
    skipped = 0

    for child in doc.element.body:
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag

        if tag == 'p':
            para = para_map.get(id(child))
            if para is None:
                continue
            para_idx += 1

            key = f"p{para_idx - 1}"
            if key in translations:
                write_to_paragraph(para, translations[key])
                written += 1
            elif para.text.strip():
                skipped += 1

        elif tag == 'tbl':
            table = table_map.get(id(child))
            if table is None:
                continue
            table_idx += 1

            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    key = f"t{table_idx - 1}_r{ri}_c{ci}"
                    if key in translations:
                        write_to_cell(cell, translations[key])
                        written += 1
                    elif cell.text.strip():
                        skipped += 1

    doc.save(output_path)

    print(f"\n{'='*50}")
    print(f"DOCX Write Complete")
    print(f"{'='*50}")
    print(f"  Written: {written} elements")
    if skipped > 0:
        print(f"  Skipped: {skipped} elements (no translation found)")
    print(f"  Output: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("Usage: python write_docx.py <input.docx> <output.docx> <translations.json>")
        sys.exit(1)

    write_docx(sys.argv[1], sys.argv[2], sys.argv[3])
